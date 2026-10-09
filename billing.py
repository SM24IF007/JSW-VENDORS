"""Admin payroll batches, reviewed attendance imports, reports and invoices."""
import csv
from contractor_profiles import CONTRACTOR_PROFILES
import json
import re
import secrets
from datetime import datetime
from decimal import Decimal, InvalidOperation, ROUND_HALF_UP
from io import BytesIO, StringIO
from zipfile import ZipFile, BadZipFile
from flask import Blueprint, request, session, redirect, url_for, render_template, flash, abort, send_file

bp = Blueprint('billing', __name__, url_prefix='/admin/billing')
D = Decimal
PROFILES = {'deva': 'Deva February reference (separate salary and OT)',
            'zinc': 'PRB Zinc March reference (wages including ex-gratia)'}
DEFAULTS = dict(days='24', hra='5', leave='4.8', bonus='8.33', pf_rate='12', pf_er='13', pf_cap='15000',
                esic_ee='0.75', esic_er='3.25', pt_high='300', pt_mid='175', holiday='0', washing='0')
CONFIG_LABELS = dict(days='Standard paid days in month', hra='HRA %', leave='Leave %', bonus='Bonus %',
                     pf_rate='Employee PF %', pf_er='Employer PF %', pf_cap='PF wage ceiling',
                     esic_ee='Employee ESIC %', esic_er='Employer ESIC %', pt_high='PT above 10,000',
                     pt_mid='PT above 7,500', holiday='Holiday encashment factor', washing='Monthly washing allowance')
TEXT_FIELDS = ('employee_code','name','designation','bank_account','ifsc','bank_address')
NUM_FIELDS = ('paid_days','ot_days','ph_days','basic_daily','da_daily','special_daily','safety_daily',
              'margin_daily','advance','food','mlwf')
FIELD_LABELS = dict(employee_code='Employee / gate-pass code', name='Employee name', designation='Designation',
    bank_account='Beneficiary account', ifsc='IFSC', bank_address='Bank branch / address',paid_days='Paid days',
    ot_days='OT days / ex-gratia units',ph_days='Paid holiday days',basic_daily='Basic per day',da_daily='DA per day',
    special_daily='Other allowance per day',safety_daily='Safety per day (cost only)',margin_daily='Service margin per day',
    advance='Advance deduction',food='Food deduction',mlwf='MLWF employee deduction')
OVERRIDE_FIELDS = ('salary_advance','ot_advance','salary_pt','ot_pt','combined_pt')
FIELD_LABELS.update(salary_advance='Salary advance (blank = combined advance)',ot_advance='OT advance (blank = zero)',salary_pt='Salary PT override (blank = automatic)',ot_pt='OT PT override (blank = automatic)',combined_pt='Combined PT override (blank = automatic)')
FIELD_LABELS['advance']='Combined advance deduction'
HEADERS = [FIELD_LABELS[k] for k in TEXT_FIELDS + NUM_FIELDS + OVERRIDE_FIELDS] + ['ESIC enabled (1/0)']


def init_billing(app, get_db, contractors, xlsx_builder):
    bp.db = get_db; bp.contractors = contractors; bp.xlsx = xlsx_builder
    c = get_db()
    c.executescript('''
    CREATE TABLE IF NOT EXISTS payroll_runs (
      id INTEGER PRIMARY KEY, contractor TEXT NOT NULL, month TEXT NOT NULL, department TEXT NOT NULL,
      profile TEXT NOT NULL, config TEXT NOT NULL, status TEXT NOT NULL DEFAULT 'Draft',
      sender_account TEXT DEFAULT '', contact TEXT DEFAULT '', cheque TEXT DEFAULT '', payment_date TEXT DEFAULT '',
      created_at TEXT DEFAULT CURRENT_TIMESTAMP, approved_at TEXT);
    CREATE TABLE IF NOT EXISTS payroll_rows (
      id INTEGER PRIMARY KEY, run_id INTEGER NOT NULL REFERENCES payroll_runs(id) ON DELETE CASCADE,
      employee_code TEXT NOT NULL, name TEXT NOT NULL, designation TEXT DEFAULT '',
      bank_account TEXT DEFAULT '', ifsc TEXT DEFAULT '', bank_address TEXT DEFAULT '',
      paid_days TEXT DEFAULT '0', ot_days TEXT DEFAULT '0', ph_days TEXT DEFAULT '0',
      basic_daily TEXT DEFAULT '0', da_daily TEXT DEFAULT '0', special_daily TEXT DEFAULT '0',
      safety_daily TEXT DEFAULT '0', margin_daily TEXT DEFAULT '0', advance TEXT DEFAULT '0',
      food TEXT DEFAULT '0', mlwf TEXT DEFAULT '0', esic_enabled INTEGER DEFAULT 1,
      UNIQUE(run_id,employee_code));
    CREATE TABLE IF NOT EXISTS payroll_imports (
      id INTEGER PRIMARY KEY, run_id INTEGER NOT NULL REFERENCES payroll_runs(id), data TEXT NOT NULL,
      warnings TEXT NOT NULL, source TEXT NOT NULL, created_at TEXT DEFAULT CURRENT_TIMESTAMP);
    CREATE TABLE IF NOT EXISTS billing_invoices (
      id INTEGER PRIMARY KEY, contractor TEXT NOT NULL, invoice_no TEXT NOT NULL,
      invoice_date TEXT NOT NULL, buyer TEXT NOT NULL, buyer_address TEXT NOT NULL,
      seller_address TEXT NOT NULL, seller_gstin TEXT DEFAULT '', buyer_gstin TEXT DEFAULT '',
      po_number TEXT DEFAULT '', bank_account TEXT DEFAULT '', ifsc TEXT DEFAULT '',
      lines TEXT NOT NULL, cgst TEXT NOT NULL, sgst TEXT NOT NULL, igst TEXT NOT NULL,
      status TEXT NOT NULL DEFAULT 'Draft', created_at TEXT DEFAULT CURRENT_TIMESTAMP,
      UNIQUE(contractor,invoice_no));
    ''')
    columns={r[1] for r in c.execute('PRAGMA table_info(payroll_rows)')}
    for key in OVERRIDE_FIELDS:
        if key not in columns:c.execute(f"ALTER TABLE payroll_rows ADD COLUMN {key} TEXT DEFAULT ''")
    invoice_columns={r[1] for r in c.execute('PRAGMA table_info(billing_invoices)')}
    if 'round_total' not in invoice_columns:c.execute("ALTER TABLE billing_invoices ADD COLUMN round_total INTEGER DEFAULT 0")
    c.commit();c.close()
    import attendance_import
    attendance_import.init_schema()
    if 'billing' not in app.blueprints: app.register_blueprint(bp)


@bp.before_request
def guard():
    if not session.get('admin_logged_in'): return redirect(url_for('login'))
    if 'billing_csrf' not in session: session['billing_csrf'] = secrets.token_urlsafe(24)
    if request.method == 'POST' and not secrets.compare_digest(request.form.get('csrf',''), session['billing_csrf']):
        abort(400, 'Your form session expired. Reload the page and try again.')


@bp.app_template_filter('money')
def money(value): return f'{cents(D(str(value))):,.2f}'


def number(value, label='Amount', maximum='1000000000'):
    try: n = D(str(value).strip() or '0')
    except (InvalidOperation, ValueError): raise ValueError(f'{label}: enter a number.')
    if not n.is_finite() or n < 0 or n > D(maximum): raise ValueError(f'{label}: enter a valid non-negative number.')
    return n


def rounded(v): return v.quantize(D('1'), rounding=ROUND_HALF_UP)
def cents(v): return v.quantize(D('.01'), rounding=ROUND_HALF_UP)
def clean(v, length=180): return str(v or '').strip()[:length]


def record(form):
    r = {k:clean(form.get(k), 120 if k=='name' else 100) for k in TEXT_FIELDS}
    if not r['employee_code'] or not r['name']: raise ValueError('Employee code and name are required.')
    r['ifsc'] = r['ifsc'].upper().replace(' ','')
    r['bank_account'] = r['bank_account'].replace(' ','')
    for k in NUM_FIELDS: r[k] = str(number(form.get(k,''), FIELD_LABELS[k], '1000000'))
    if D(r['paid_days']) > 31 or D(r['ph_days']) > 31 or D(r['ot_days']) > 100:
        raise ValueError('Check paid days, holiday days and OT units; values exceed the supported range.')
    for k in OVERRIDE_FIELDS:r[k]=str(number(form.get(k),FIELD_LABELS[k])) if str(form.get(k,'')).strip() else ''
    r['esic_enabled'] = 0 if str(form.get('esic_enabled','1')).strip().lower() in ('0','no','false') else 1
    return r


def calculate(row, config, profile, mode='combined'):
    r={k:number(row[k]) for k in NUM_FIELDS}; c={k:number(v) for k,v in config.items()}
    paid=r['paid_days'] if mode!='ot' else D(0)
    units=r['ot_days'] if mode!='salary' else D(0)
    ph=r['ph_days'] if mode!='ot' else D(0)
    advance=r['advance'] if mode!='ot' else D(0); food=r['food'] if mode!='ot' else D(0)
    if profile=='deva' and mode in ('salary','ot'):
        override=row.get(mode+'_advance','')
        if override!='':advance=number(override)
    mlwf=r['mlwf'] if mode!='ot' else D(0)
    daily=r['basic_daily']+r['da_daily']
    basic=rounded(paid*r['basic_daily']) if profile=='deva' else paid*r['basic_daily']
    da=paid*r['da_daily']; normal=basic+da
    hra=paid*daily*c['hra']/100 if profile=='deva' else normal*c['hra']/100
    ot=daily*units; holiday=rounded(daily*ph) if profile=='deva' else normal*c['holiday']
    extra=rounded(paid*r['special_daily'])
    leave=rounded(paid*rounded(daily*c['leave']/100)) if profile=='deva' else normal*c['leave']/100
    bonus=rounded(paid*rounded(daily*c['bonus']/100)) if profile=='deva' else normal*c['bonus']/100
    washing=D(0) if profile=='deva' else c['washing']/c['days']*paid
    safety=(paid+units)*r['safety_daily']; margin=(paid+units)*r['margin_daily']
    gross=normal+hra+ot+holiday+extra+leave+bonus+washing
    pf_base=normal if profile=='deva' else normal+holiday+leave
    pf=min(pf_base,c['pf_cap'])*c['pf_rate']/100
    pf_er=min(pf_base,c['pf_cap'])*c['pf_er']/100
    esi_base=gross if profile=='deva' else gross-washing
    esi_active=bool(row['esic_enabled']) and (profile=='deva' or esi_base<=21000)
    esic=esi_base*c['esic_ee']/100 if esi_active else D(0)
    er_base=normal+hra+holiday+ot+extra if profile=='deva' else esi_base
    esic_er=er_base*c['esic_er']/100 if esi_active else D(0)
    pt=(c['pt_high'] if gross>0 else D(0)) if profile=='zinc' else (c['pt_high'] if gross>10000 else c['pt_mid'] if gross>7500 else D(0))
    if profile=='deva' and row.get(mode+'_pt','')!='':pt=number(row[mode+'_pt'])
    deductions=pf+esic+pt+mlwf+advance+food
    if profile=='deva': deductions=rounded(deductions)
    net=gross-deductions
    if profile=='zinc': net=rounded(net)
    cost=gross+pf_er+esic_er+safety+margin+mlwf*3
    return dict(basic=basic,da=da,hra=hra,ot=ot,holiday=holiday,extra=extra,leave=leave,bonus=bonus,
        washing=washing,gross=gross,pf=pf,esic=esic,pt=pt,mlwf=mlwf,advance=advance,food=food,
        deductions=deductions,net=net,pf_er=pf_er,esic_er=esic_er,safety=safety,margin=margin,cost=cost)


def load_run(rid):
    c=bp.db();run=c.execute('SELECT * FROM payroll_runs WHERE id=?',(rid,)).fetchone()
    rows=c.execute('SELECT * FROM payroll_rows WHERE run_id=? ORDER BY name,id',(rid,)).fetchall();c.close()
    if run is None: abort(404)
    return dict(run),[dict(r) for r in rows]


def draft(run):
    if run['status']!='Draft': raise ValueError('This batch is approved and locked. Reopen it before editing.')


def save_row(c,rid,r,rowid=None):
    keys=list(TEXT_FIELDS+NUM_FIELDS+OVERRIDE_FIELDS)+['esic_enabled']
    if rowid:
        c.execute('UPDATE payroll_rows SET '+','.join(k+'=?' for k in keys)+' WHERE id=? AND run_id=?',
                  [r[k] for k in keys]+[rowid,rid])
    else:
        c.execute('INSERT INTO payroll_rows (run_id,'+','.join(keys)+') VALUES ('+','.join('?' for _ in range(len(keys)+1))+')',
                  [rid]+[r[k] for k in keys])


def review(run,rows):
    messages=[]; config=json.loads(run['config']);accounts={}
    if not rows: messages.append('Add or import employees before approval.')
    for row in rows:
        if D(row['paid_days'])>D(config['days']): messages.append(f"{row['name']}: paid days exceed the batch standard days.")
        if D(row['basic_daily'])+D(row['da_daily'])<=0: messages.append(f"{row['name']}: basic and DA rates are missing.")
        modes=['salary','ot'] if run['profile']=='deva' else ['combined']
        for mode in modes:
            if calculate(row,config,run['profile'],mode)['net']<0: messages.append(f"{row['name']}: {mode} net pay is negative.")
        account=row['bank_account']
        if account:
            if account in accounts: messages.append(f"Duplicate bank account: {accounts[account]} and {row['name']}.")
            accounts[account]=row['name']
    return messages


@bp.route('/')
def home():
    c=bp.db();runs=c.execute('SELECT * FROM payroll_runs ORDER BY id DESC').fetchall()
    invoices=c.execute('SELECT * FROM billing_invoices ORDER BY id DESC').fetchall();c.close()
    return render_template('billing/home.html',runs=runs,invoices=invoices,contractors=bp.contractors,profiles=PROFILES)


@bp.route('/create',methods=['POST'])
def create():
    try:
        firm=request.form.get('contractor');profile=request.form.get('profile');month=request.form.get('month','')
        if firm not in bp.contractors or profile not in PROFILES: raise ValueError('Choose a contractor and calculation profile.')
        datetime.strptime(month,'%Y-%m');department=clean(request.form.get('department'))
        if not department: raise ValueError('Department / work area is required.')
        config=DEFAULTS.copy()
        if profile=='zinc': config.update(days='26',leave='5.77',holiday=str(D(8)/12/26),washing='1500',pt_high='200')
        c=bp.db();cur=c.execute('INSERT INTO payroll_runs (contractor,month,department,profile,config) VALUES (?,?,?,?,?)',
                                (firm,month,department,profile,json.dumps(config)));rid=cur.lastrowid;c.commit();c.close()
        return redirect(url_for('billing.run',rid=rid))
    except ValueError as e: flash(str(e),'error');return redirect(url_for('billing.home'))


@bp.route('/runs/<int:rid>')
def run(rid):
    batch,rows=load_run(rid);config=json.loads(batch['config']);mode=request.args.get('mode','salary' if batch['profile']=='deva' else 'combined')
    if mode not in ('salary','ot','combined') or (batch['profile']=='zinc' and mode!='combined'):mode='combined'
    results=[dict(row=row,calc=calculate(row,config,batch['profile'],mode)) for row in rows]
    totals={key:sum((cents(r['calc'][key]) for r in results),D(0)) for key in ('gross','deductions','net','cost')}
    issues=review(batch,rows)
    query=clean(request.args.get('q')).lower()
    matching=[r for r in results if query in (r['row']['name']+' '+r['row']['employee_code']).lower()]
    pages=max(1,(len(matching)+14)//15);page=max(1,min(request.args.get('page',1,type=int),pages))
    visible=matching[(page-1)*15:page*15]
    c=bp.db();attachments=c.execute('SELECT id,filename,created_at FROM attendance_uploads WHERE run_id=? ORDER BY id DESC',(rid,)).fetchall();c.close()
    return render_template('billing/run.html',batch=batch,results=visible,employee_count=len(rows),query=query,page=page,pages=pages,matched_count=len(matching),attachments=attachments,totals=totals,config=config,
        config_labels=CONFIG_LABELS,fields=FIELD_LABELS,text_fields=TEXT_FIELDS,num_fields=NUM_FIELDS,override_fields=OVERRIDE_FIELDS,
        profiles=PROFILES,issues=issues,mode=mode)


@bp.route('/runs/<int:rid>/settings',methods=['POST'])
def settings(rid):
    run,rows=load_run(rid)
    try:
        draft(run);config={k:str(number(request.form.get(k),v)) for k,v in CONFIG_LABELS.items()}
        if not 1<=D(config['days'])<=31:raise ValueError('Standard days must be between 1 and 31.')
        for key in ('hra','leave','bonus','pf_rate','pf_er','esic_ee','esic_er'):
            if D(config[key])>100:raise ValueError('Percentage rates cannot exceed 100.')
        paydate=clean(request.form.get('payment_date'))
        if paydate:datetime.strptime(paydate,'%Y-%m-%d')
        c=bp.db();c.execute('UPDATE payroll_runs SET config=?,sender_account=?,contact=?,cheque=?,payment_date=? WHERE id=?',
            (json.dumps(config),clean(request.form.get('sender_account')),clean(request.form.get('contact')),
             clean(request.form.get('cheque')),paydate,rid));c.commit();c.close();flash('Calculation and payment settings saved.','success')
    except ValueError as e:flash(str(e),'error')
    return redirect(url_for('billing.run',rid=rid))


@bp.route('/runs/<int:rid>/row',methods=['POST'])
def row_save(rid):
    run,rows=load_run(rid)
    try:
        draft(run);r=record(request.form);rowid=request.form.get('row_id',type=int)
        if rowid and not any(x['id']==rowid for x in rows):abort(404)
        c=bp.db()
        try:save_row(c,rid,r,rowid);c.commit()
        except Exception:
            c.rollback();raise ValueError('Employee code already exists in this batch. Edit that row instead.')
        finally:c.close()
        flash('Employee saved and wages recalculated.','success')
    except ValueError as e:flash(str(e),'error')
    return redirect(url_for('billing.run',rid=rid))


@bp.route('/runs/<int:rid>/row/<int:rowid>/delete',methods=['POST'])
def row_delete(rid,rowid):
    run,_=load_run(rid)
    try:
        draft(run);c=bp.db();c.execute('DELETE FROM payroll_rows WHERE id=? AND run_id=?',(rowid,rid));c.commit();c.close()
    except ValueError as e:flash(str(e),'error')
    return redirect(url_for('billing.run',rid=rid))


@bp.route('/runs/<int:rid>/status',methods=['POST'])
def status(rid):
    run,rows=load_run(rid)
    try:
        action=request.form.get('action')
        if action=='approve':
            problems=review(run,rows)
            if problems:raise ValueError('Resolve review issues before approval: '+problems[0])
            if request.form.get('confirmed')!='yes':raise ValueError('Confirm that rates, deductions and salary/OT separation were reviewed.')
            value='Approved'
        elif action=='reopen':value='Draft'
        else:abort(400)
        c=bp.db();c.execute('UPDATE payroll_runs SET status=?,approved_at=? WHERE id=?',
            (value,datetime.now().isoformat() if value=='Approved' else None,rid));c.commit();c.close()
        flash('Batch '+value.lower()+'.','success')
    except ValueError as e:flash(str(e),'error')
    return redirect(url_for('billing.run',rid=rid))


def read_import(upload,sheet_name,profile):
    """Read the provided workbook layouts or a simple attendance-summary template."""
    filename=upload.filename or '';data=upload.read();warnings=[];split_values={}
    if filename.lower().endswith('.csv'):
        text=data.decode('utf-8-sig');grid=list(csv.reader(StringIO(text)));chosen='CSV'
    elif filename.lower().endswith('.xlsx'):
        try:
            with ZipFile(BytesIO(data)) as z:
                if sum(x.file_size for x in z.infolist())>50*1024*1024:raise ValueError('Expanded workbook exceeds 50 MB.')
            import openpyxl
            w=openpyxl.load_workbook(BytesIO(data),data_only=True,read_only=True)
            if sheet_name and sheet_name not in w.sheetnames:raise ValueError('Sheet not found. Available sheets: '+', '.join(w.sheetnames))
            s=w[sheet_name] if sheet_name else w.worksheets[0];chosen=s.title
            if s.max_row>5000 or s.max_column>200:raise ValueError('Use a workbook with at most 5,000 rows and 200 columns.')
            grid=list(s.iter_rows(values_only=True))
            if profile=='deva':
                for ws in w:
                    mode={'26 DAYS':'salary','OVER TIME':'ot'}.get(ws.title.strip().upper())
                    if not mode:continue
                    if ws.max_row>5000 or ws.max_column>200:continue
                    for vals in ws.iter_rows(min_row=12,values_only=True):
                        if len(vals)<44 or not isinstance(vals[0],(int,float)) or not vals[5]:continue
                        split_values.setdefault(str(vals[5]).strip(),{}).update({mode+'_advance':vals[43] or 0,mode+'_pt':vals[41] or 0})
            w.close()
        except (ValueError,ImportError):raise
        except Exception:raise ValueError('The uploaded file is not a readable XLSX workbook. Re-save it as XLSX and try again.')
    else:raise ValueError('Upload an XLSX or UTF-8 CSV file.')
    if len(grid)>5000:raise ValueError('Import limit is 5,000 rows.')
    norm=lambda v:re.sub(r'[^a-z0-9]','',str(v or '').lower())
    header=None;layout='generic'
    for index,row in enumerate(grid[:30]):
        labels=[norm(x) for x in row]
        if 'nameofworkmen' in labels and 'paiddays' in labels:header=index;layout='deva';break
        if 'employeename' in labels and 'totaldayspresent' in labels:header=index;layout='zinc';break
        if 'employeegatepasscode' in labels or ('employeecode' in labels and 'name' in labels):header=index;break
    if header is None:raise ValueError('No supported attendance-summary header found. Download the import template. Daily P/A calendars need a confirmed mapping first.')
    if layout=='deva' and profile!='deva' or layout=='zinc' and profile!='zinc':raise ValueError('The workbook layout does not match this batch calculation profile.')
    aliases={norm(v):k for k,v in FIELD_LABELS.items()}
    aliases.update({norm(k):k for k in TEXT_FIELDS+NUM_FIELDS+OVERRIDE_FIELDS});aliases.update({'esicenabled10':'esic_enabled','esicenabled':'esic_enabled','name':'name','employeecode':'employee_code','otdays':'ot_days'})
    mappings={i:aliases[norm(x)] for i,x in enumerate(grid[header]) if norm(x) in aliases}
    records=[];seen=set()
    for rn,values in enumerate(grid[header+1:],header+2):
        def v(i):return values[i-1] if len(values)>=i and values[i-1] is not None else ''
        def val(x):
            if isinstance(x,float) and x.is_integer():return str(int(x))
            return str(x or '') if x!=0 else '0'
        if layout=='deva':
            if not v(7) or not isinstance(v(1),(int,float)):continue
            if not v(25) and not v(28):continue
            r=dict(employee_code=val(v(6)),name=v(7),designation=v(8),paid_days=v(25),ot_days=v(28),ph_days=v(27),
                basic_daily=v(17),da_daily=v(18),special_daily=v(20),safety_daily=v(21),margin_daily=v(49),advance=v(44),mlwf=v(43))
            r['combined_pt']=v(42)
            r.update(split_values.get(r['employee_code'],{}))
        elif layout=='zinc':
            if not v(4) or not isinstance(v(1),(int,float)):continue
            r=dict(employee_code=val(v(3)),name=v(4),designation=v(8),paid_days=v(13),ot_days=v(14),basic_daily=v(16),da_daily=v(17),
                advance=v(35),food=v(34),mlwf=v(30),bank_account=val(v(38)),ifsc=val(v(39)),esic_enabled='0' if v(31)==0 and v(32)==0 else '1')
            if r['bank_account'] and v(39) and not re.fullmatch(r'[A-Z]{4}0[A-Z0-9]{6}',val(v(39)).upper()) and norm(v(39))!=norm(r['name']):
                r['bank_account']='';warnings.append(f'Row {rn}: bank beneficiary text does not match employee name; account withheld until manually verified.')
            if not re.fullmatch(r'[A-Z]{4}0[A-Z0-9]{6}',r['ifsc'].upper()):
                matches=[val(x) for x in values[38:] if re.fullmatch(r'[A-Z]{4}0[A-Z0-9]{6}',val(x).upper())]
                if len(matches)==1:r['ifsc']=matches[0];warnings.append(f'Row {rn}: IFSC found after the labelled column; verify the bank details.')
        else:
            if not any(x is not None and str(x).strip() for x in values):continue
            r={key:val(values[i]) if i<len(values) else '' for i,key in mappings.items()}
        if not r.get('employee_code'):
            r['employee_code']=f'IMPORT-{rn}';warnings.append(f'Row {rn}: employee code missing; temporary code assigned for review.')
        for key in ('basic_daily','da_daily','paid_days','ot_days'):
            if r.get(key) in (None,''):warnings.append(f'Row {rn}: {FIELD_LABELS[key]} blank; imported as zero.')
        try:r=record(r)
        except ValueError as e:raise ValueError(f'Row {rn}: {e}')
        if r['employee_code'] in seen:raise ValueError(f"Duplicate employee code in import: {r['employee_code']}.")
        seen.add(r['employee_code']);records.append(r)
    if not records:raise ValueError('No employee rows found. Formula cells need saved Excel values; open, recalculate and save the workbook if needed.')
    if layout=='deva' and split_values:warnings.insert(0,'Imported salary/OT PT and advance overrides from the separate sheets. These are fixed amounts; clear an override to calculate PT automatically after changing attendance or rates.')
    return records,warnings,chosen


@bp.route('/runs/<int:rid>/import',methods=['POST'])
def import_preview(rid):
    batch,_=load_run(rid)
    try:
        draft(batch);upload=request.files.get('file')
        if upload is None or not upload.filename:raise ValueError('Choose a workbook or CSV file.')
        rows,warnings,sheet=read_import(upload,clean(request.form.get('sheet')),batch['profile'])
        c=bp.db();cur=c.execute('INSERT INTO payroll_imports (run_id,data,warnings,source) VALUES (?,?,?,?)',
            (rid,json.dumps(rows),json.dumps(warnings),clean(upload.filename)+' / '+sheet));iid=cur.lastrowid;c.commit();c.close()
        return redirect(url_for('billing.preview',iid=iid))
    except (ValueError,UnicodeDecodeError,ImportError) as e:
        flash(str(e) if not isinstance(e,ImportError) else 'Install requirements.txt to enable Excel import.','error')
        return redirect(url_for('billing.run',rid=rid))


@bp.route('/imports/<int:iid>')
def preview(iid):
    c=bp.db();item=c.execute('SELECT * FROM payroll_imports WHERE id=?',(iid,)).fetchone();c.close()
    if not item:abort(404)
    batch,existing=load_run(item['run_id']);rows=json.loads(item['data']);codes={x['employee_code'] for x in existing}
    conflicts=[r['employee_code'] for r in rows if r['employee_code'] in codes]
    return render_template('billing/preview.html',item=item,batch=batch,rows=rows,warnings=json.loads(item['warnings']),conflicts=conflicts)


@bp.route('/imports/<int:iid>/confirm',methods=['POST'])
def confirm_import(iid):
    c=bp.db();item=c.execute('SELECT * FROM payroll_imports WHERE id=?',(iid,)).fetchone();c.close()
    if not item:abort(404)
    batch,_=load_run(item['run_id'])
    try:
        draft(batch)
        if request.form.get('reviewed')!='yes':raise ValueError('Confirm the import review first.')
        c=bp.db()
        try:
            for row in json.loads(item['data']):save_row(c,batch['id'],row)
            c.execute('DELETE FROM payroll_imports WHERE id=?',(iid,));c.commit()
        except Exception:
            c.rollback();raise ValueError('Import not applied: an employee code already exists. Use a new batch or edit existing rows.')
        finally:c.close()
        flash('Attendance imported. Review rates and bank details before approval.','success')
    except ValueError as e:flash(str(e),'error')
    return redirect(url_for('billing.run',rid=batch['id']))


@bp.route('/template.xlsx')
def template():
    # A real file avoids per-request workbook generation and BytesIO file-wrapper failures.
    from pathlib import Path
    path=Path(__file__).parent/'resources'/'Attendance_Import_Template.xlsx'
    if not path.is_file():
        flash('The attendance template is missing. Extract the complete updated ZIP, including resources.','error')
        return redirect(url_for('billing.home'))
    return send_file(path,as_attachment=True,download_name='Attendance_Import_Template.xlsx',
        mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',conditional=True)



REPORT_KEYS=('basic','da','hra','holiday','ot','extra','leave','bonus','washing','gross','pf','esic','pt','mlwf','advance','food','deductions','net','pf_er','esic_er','safety','margin','cost')
REPORT_LABELS=('Basic','DA','HRA','Holiday','OT / ex-gratia','Other','Leave','Bonus','Washing','Gross','PF EE','ESIC EE','PT','MLWF','Advance','Food','Deductions','Net pay','PF ER','ESIC ER','Safety cost','Service margin','Total cost')


@bp.route('/runs/<int:rid>/export/<kind>.<fmt>')
def export(rid,kind,fmt):
    batch,rows=load_run(rid)
    if kind not in ('wages','bank') or fmt not in ('pdf','xlsx'):abort(404)
    mode=request.args.get('mode','salary')
    if mode not in ('salary','ot','combined'):abort(400)
    if batch['profile']=='zinc' and mode!='combined':
        flash('The Zinc reference keeps ex-gratia inside wages. Use Combined; a separate OT split is not defined.','error')
        return redirect(url_for('billing.run',rid=rid))
    config=json.loads(batch['config']);items=[(r,calculate(r,config,batch['profile'],mode)) for r in rows]
    if kind=='bank':
        issues=[]
        if batch['status']!='Approved':issues.append('Approve the reviewed batch before generating a bank-payment list.')
        if not re.fullmatch(r'\d{6,30}',batch['sender_account']):issues.append('Enter the company debit account in payment settings.')
        if not batch['payment_date']:issues.append('Enter a payment date.')
        items=[(r,v) for r,v in items if cents(v['net'])!=0]
        seen=set()
        for r,v in items:
            if v['net']<0 or not re.fullmatch(r'\d{6,30}',r['bank_account']) or not re.fullmatch(r'[A-Z]{4}0[A-Z0-9]{6}',r['ifsc']):issues.append(r['name']+': check net pay, bank account and IFSC.')
            if r['bank_account'] in seen:issues.append('Duplicate beneficiary account.')
            seen.add(r['bank_account'])
        if not items:issues.append('No non-zero payments to export.')
        if issues:flash(' '.join(issues[:5]),'error');return redirect(url_for('billing.run',rid=rid,mode=mode))
        headers=['SR.NO','AMT','ACCOUNT NO','BENEF ACCOUNT NO','BEN NAME','BEN ADDRESS','SENDER','IFSC CODE','CONTACT']
        data=[[i,str(cents(v['net'])),batch['sender_account'],r['bank_account'],r['name'],r['bank_address'],batch['contractor'],r['ifsc'],batch['contact']] for i,(r,v) in enumerate(items,1)]
        total=sum((cents(v['net']) for r,v in items),D(0));data.append(['TOTAL',str(total)]+['']*7)
    else:
        headers=['Code','Employee','Paid days','OT units']+list(REPORT_LABELS)
        data=[[r['employee_code'],r['name'],r['paid_days'] if mode!='ot' else '0',r['ot_days'] if mode!='salary' else '0']+[str(cents(v[k])) for k in REPORT_KEYS] for r,v in items]
        data.append(['TOTAL','','','']+[str(sum((cents(v[k]) for _,v in items),D(0))) for k in REPORT_KEYS])
    title=f"{batch['contractor']} - {batch['month']} - {batch['department']} - {mode.upper()} {kind.upper()}"
    if fmt=='xlsx':out=bp.xlsx(headers,data,'Bank Payment' if kind=='bank' else 'Wages')
    else:
        from billing_pdf import report_pdf
        out=report_pdf(batch,mode,kind,headers,data,items)
    return send_file(out,as_attachment=True,download_name=f"{kind}_{mode}_{batch['month']}_{rid}.{fmt}",
        mimetype='application/pdf' if fmt=='pdf' else 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')


def invoice_amounts(invoice):
    lines=json.loads(invoice['lines']);subtotal=sum((cents(D(x['quantity'])*D(x['rate'])) for x in lines),D(0))
    taxes={k:cents(subtotal*D(invoice[k])/100) for k in ('cgst','sgst','igst')}
    total=subtotal+sum(taxes.values(),D(0))
    if dict(invoice).get('round_total'):total=rounded(total)
    return subtotal,taxes,total


@bp.route('/invoices/new',methods=['GET','POST'])
@bp.route('/invoices/<int:iid>/edit',methods=['GET','POST'])
def invoice_new(iid=None):
    existing={}
    if iid:
        c=bp.db(); item=c.execute('SELECT * FROM billing_invoices WHERE id=?',(iid,)).fetchone();c.close()
        if not item:abort(404)
        existing=dict(item)
        if existing['status']!='Draft':
            flash('Issued invoices are locked.','error');return redirect(url_for('billing.invoice_view',iid=iid))
    initial=dict(existing)
    for i,line in enumerate(json.loads(existing.get('lines','[]')),1):
        for key,value in line.items():initial[f'{key}_{i}']=value
    if request.method=='POST':
        try:
            form=request.form
            if form.get('contractor') not in bp.contractors:raise ValueError('Choose a contractor.')
            values={k:clean(form.get(k),500 if 'address' in k else 180) for k in ('contractor','invoice_no','invoice_date','buyer','buyer_address','seller_address','seller_gstin','buyer_gstin','po_number','bank_account','ifsc')}
            values['seller_address']=CONTRACTOR_PROFILES[values['contractor']]['address']
            if not all(values[k] for k in ('invoice_no','invoice_date','buyer','buyer_address','seller_address')):raise ValueError('Complete the invoice number, date, seller and buyer details.')
            datetime.strptime(values['invoice_date'],'%Y-%m-%d')
            lines=[]
            for i in range(1,7):
                desc=clean(form.get(f'description_{i}'),500)
                if not desc:continue
                qty=number(form.get(f'quantity_{i}'),'Quantity');rate=number(form.get(f'rate_{i}'),'Rate')
                if qty<=0:raise ValueError('Line quantities must be greater than zero.')
                lines.append(dict(description=desc,sac=clean(form.get(f'sac_{i}'),20),quantity=str(qty),rate=str(rate)))
            if not lines:raise ValueError('Add at least one work description and rate.')
            values['lines']=json.dumps(lines)
            values['round_total']=1 if form.get('round_total')=='yes' else 0
            for key in ('cgst','sgst','igst'):
                values[key]=str(number(form.get(key),key.upper(),'100'))
            if D(values['igst']) and (D(values['cgst']) or D(values['sgst'])):raise ValueError('Use IGST or CGST/SGST, not both.')
            c=bp.db()
            try:
                if iid:
                    c.execute('UPDATE billing_invoices SET '+','.join(k+'=?' for k in values)+' WHERE id=?',list(values.values())+[iid])
                else:
                    cur=c.execute('INSERT INTO billing_invoices ('+','.join(values)+') VALUES ('+','.join('?' for _ in values)+')',list(values.values()));iid=cur.lastrowid
                c.commit()
            except Exception:
                c.rollback();raise ValueError('This invoice number already exists for the contractor.')
            finally:c.close()
            return redirect(url_for('billing.invoice_view',iid=iid))
        except ValueError as e:flash(str(e),'error')
    return render_template('billing/invoice_new.html',contractors=bp.contractors,values=request.form if request.method=='POST' else initial,iid=iid,contractor_profiles=CONTRACTOR_PROFILES)


@bp.route('/invoices/<int:iid>')
def invoice_view(iid):
    c=bp.db();inv=c.execute('SELECT * FROM billing_invoices WHERE id=?',(iid,)).fetchone();c.close()
    if not inv:abort(404)
    subtotal,taxes,total=invoice_amounts(inv)
    lines=json.loads(inv['lines'])
    for line in lines:line['amount']=cents(D(line['quantity'])*D(line['rate']))
    return render_template('billing/invoice_view.html',invoice=inv,profile=CONTRACTOR_PROFILES[inv['contractor']],lines=lines,subtotal=subtotal,taxes=taxes,total=total,round_off=total-subtotal-sum(taxes.values(),D(0)))


@bp.route('/invoices/<int:iid>/issue',methods=['POST'])
def invoice_issue(iid):
    if request.form.get('confirmed')!='yes':flash('Confirm the invoice details before issuing.','error')
    else:
        c=bp.db();c.execute("UPDATE billing_invoices SET status='Issued' WHERE id=?",(iid,));c.commit();c.close()
    return redirect(url_for('billing.invoice_view',iid=iid))


@bp.route('/invoices/<int:iid>.pdf')
def invoice_pdf(iid):
    c=bp.db();inv=c.execute('SELECT * FROM billing_invoices WHERE id=?',(iid,)).fetchone();c.close()
    if not inv:abort(404)
    from billing_pdf import make_invoice
    return send_file(make_invoice(inv,invoice_amounts(inv)),as_attachment=True,download_name=f'invoice_{iid}.pdf',mimetype='application/pdf')
