"""Reviewed monthly attachments and explicit column/day-code mapping."""
import csv,json,re
from io import BytesIO,StringIO
from zipfile import ZipFile
from flask import request,render_template,redirect,url_for,flash,abort,send_file
from billing import bp,load_run,draft,record,number,clean,TEXT_FIELDS,NUM_FIELDS,OVERRIDE_FIELDS,FIELD_LABELS


def init_schema():
    c=bp.db();c.execute('''CREATE TABLE IF NOT EXISTS attendance_uploads (
      id INTEGER PRIMARY KEY,run_id INTEGER NOT NULL,filename TEXT NOT NULL,content BLOB NOT NULL,
      created_at TEXT DEFAULT CURRENT_TIMESTAMP)''');c.commit();c.close()


def grid_read(data,filename,sheet=''):
    if filename.lower().endswith('.csv'):
        grid=list(csv.reader(StringIO(data.decode('utf-8-sig'))));names=['CSV'];selected='CSV'
    else:
        try:
            with ZipFile(BytesIO(data)) as z:
                if sum(x.file_size for x in z.infolist())>50*1024*1024:raise ValueError('Expanded workbook exceeds 50 MB.')
            import openpyxl
            w=openpyxl.load_workbook(BytesIO(data),read_only=True,data_only=True)
            names=w.sheetnames
            if sheet and sheet not in names:raise ValueError('Select a worksheet from this workbook.')
            ws=w[sheet] if sheet else w.worksheets[0];selected=ws.title
            if (ws.max_row or 0)>5000 or (ws.max_column or 0)>200:raise ValueError('Maximum 5,000 rows and 200 columns per sheet.')
            grid=list(ws.iter_rows(values_only=True));w.close()
        except ValueError:raise
        except Exception:raise ValueError('Cannot read this workbook. Re-save it as XLSX and try again.')
    if len(grid)>5000 or any(len(r)>200 for r in grid):raise ValueError('Maximum 5,000 rows and 200 columns per sheet.')
    grid=[[str(v) if v is not None else '' for v in row] for row in grid]
    if not grid:raise ValueError('The selected sheet is empty.')
    return grid,names,selected


def load_upload(uid):
    c=bp.db();u=c.execute('SELECT * FROM attendance_uploads WHERE id=?',(uid,)).fetchone();c.close()
    if not u:abort(404)
    batch,_=load_run(u['run_id']);return dict(u),batch


@bp.route('/runs/<int:rid>/attendance',methods=['POST'])
def attendance_upload(rid):
    batch,_=load_run(rid)
    try:
        draft(batch);file=request.files.get('file')
        if not file or not file.filename:raise ValueError('Attach the attendance sheet for this month.')
        filename=clean(file.filename.replace('\\','/').split('/')[-1],150);data=file.read()
        ext=filename.rsplit('.',1)[-1].lower()
        if ext not in ('xlsx','csv','pdf','png','jpg','jpeg'):raise ValueError('Attach XLSX, CSV, PDF, JPG or PNG.')
        if ext in ('xlsx','csv'):grid_read(data,filename)
        elif ext=='pdf' and not data.startswith(b'%PDF-'):raise ValueError('Invalid PDF attachment.')
        elif ext in ('png','jpg','jpeg'):
            from PIL import Image
            try:Image.open(BytesIO(data)).verify()
            except Exception:raise ValueError('Invalid image attachment.')
        c=bp.db();cur=c.execute('INSERT INTO attendance_uploads(run_id,filename,content) VALUES(?,?,?)',(rid,filename,data));uid=cur.lastrowid;c.commit();c.close()
        if ext in ('xlsx','csv'):return redirect(url_for('billing.attendance_map',uid=uid))
        flash('Attendance attached. For this PDF/photo, enter checked totals manually or upload its Excel/CSV version.','success')
    except (ValueError,UnicodeError) as e:flash(str(e),'error')
    return redirect(url_for('billing.run',rid=rid))


@bp.route('/attendance/<int:uid>/download')
def attendance_download(uid):
    u,_=load_upload(uid)
    return send_file(BytesIO(u['content']),as_attachment=True,download_name=u['filename'],mimetype='application/octet-stream')


def map_records(grid,form):
    """No inferred attendance semantics: mappings and day weights are explicit."""
    try:header=int(form.get('header','1'));last=int(form.get('last_row') or len(grid))
    except ValueError:raise ValueError('Enter valid header and last data row numbers.')
    if not 1<=header<last<=len(grid):raise ValueError('Select a header row followed by the data rows.')
    width=max(map(len,grid));mapping={}
    for key in TEXT_FIELDS+NUM_FIELDS+OVERRIDE_FIELDS:
        raw=form.get('col_'+key,'')
        if raw!='':
            try:col=int(raw)
            except ValueError:raise ValueError('Invalid column selection.')
            if not 0<=col<width:raise ValueError('Selected column is outside the sheet.')
            mapping[key]=col
    if 'name' not in mapping or 'employee_code' not in mapping:raise ValueError('Map employee name and employee code columns.')
    daily=form.get('attendance_mode')=='daily';weights={}
    if daily:
        try:start=int(form.get('day_start',''));end=int(form.get('day_end',''))
        except ValueError:raise ValueError('Select the first and last daily attendance columns.')
        if not 0<=start<=end<width or end-start>=31:raise ValueError('Select between 1 and 31 consecutive daily attendance columns.')
        for item in form.get('day_codes','').split(','):
            pair=item.strip().split('=')
            if len(pair)!=2 or not pair[0].strip():raise ValueError('Enter attendance weights as P=1,A=0,HD=0.5.')
            key=pair[0].strip().upper()
            if key in weights:raise ValueError('A day code is repeated.')
            weights[key]=number(pair[1],'Day weight','1')
        if form.get('rules_confirmed')!='yes':raise ValueError('Confirm the day-code meanings and blank-cell rule.')
    elif 'paid_days' not in mapping:raise ValueError('Map the total paid-days column or choose daily attendance codes.')
    divisor=number(form.get('hours_per_day',''),'Hours per day','24')
    hours=form.get('ot_unit')=='hours'
    if hours and divisor<=0:raise ValueError('Enter the agreed hours per OT day to convert hours.')
    factor=number(form.get('ot_factor','1'),'OT multiplier','10')
    if factor<=0:raise ValueError('OT multiplier must be greater than zero.')
    if form.get('units_confirmed')!='yes':raise ValueError('Confirm whether OT uses hours or day-equivalent units and the agreed multiplier.')
    default_basic=str(number(form.get('default_basic','0'),'Default basic per day','1000000'))
    default_da=str(number(form.get('default_da','0'),'Default DA per day','1000000'))
    rows=[];warnings=[];seen=set()
    def at(values,index):return values[index].strip() if index<len(values) else ''
    for n,values in enumerate(grid[header:last],header+1):
        if not any(v.strip() for v in values):continue
        r={key:at(values,col) for key,col in mapping.items()}
        if not r.get('name') and not r.get('employee_code'):continue
        if r.get('name','').strip().lower() in ('total','grand total','subtotal') or r.get('employee_code','').strip().lower() in ('total','grand total','subtotal'):continue
        if not r.get('basic_daily'):r['basic_daily']=default_basic
        if not r.get('da_daily'):r['da_daily']=default_da
        if daily:
            paid=number(0)
            for col in range(start,end+1):
                code=at(values,col).upper()
                if not code and form.get('blank_rule')=='zero':continue
                if code not in weights:raise ValueError(f'Row {n}, column {col+1}: unknown attendance code {code or "(blank)"}. Define it or correct the sheet.')
                paid+=weights[code]
            r['paid_days']=str(paid)
        ot=number(r.get('ot_days','0'),'OT units')
        r['ot_days']=str((ot/divisor if hours else ot)*factor)
        try:r=record(r)
        except ValueError as e:raise ValueError(f'Row {n}: {e}')
        if r['employee_code'] in seen:raise ValueError(f'Row {n}: duplicate employee code {r["employee_code"]}.')
        seen.add(r['employee_code']);rows.append(r)
    if not rows:raise ValueError('No employee rows found in the chosen data range.')
    if any(number(r['basic_daily'])+number(r['da_daily'])==0 for r in rows):warnings.append('Some daily wage rates are zero. Fill them before approval.')
    warnings.append('Check bank accounts against the source; Excel numeric cells can lose leading zeros or digits.')
    warnings.append(f'OT conversion: {"hours / "+str(divisor) if hours else "day-equivalent units"} × {factor}.')
    return rows,warnings


@bp.route('/attendance/<int:uid>/map',methods=['GET','POST'])
def attendance_map(uid):
    u,batch=load_upload(uid)
    try:
        draft(batch)
        if not u['filename'].lower().endswith(('.csv','.xlsx')):raise ValueError('Use Excel/CSV to map attendance columns.')
        form=request.form if request.method=='POST' else request.args
        grid,names,sheet=grid_read(u['content'],u['filename'],form.get('sheet',''))
        try:header=int(form.get('header','1'))
        except ValueError:header=1
        header=max(1,min(header,len(grid)));width=max(map(len,grid))
        if request.method=='POST':
            try:rows,warnings=map_records(grid,form)
            except ValueError as e:
                flash(str(e),'error')
                return render_template('billing/map.html',batch=batch,upload=u,sheets=names,sheet=sheet,header=header,columns=[(i,str(grid[header-1][i] if i<len(grid[header-1]) else '')[:80]) for i in range(width)],sample=list(enumerate(grid[max(0,header-2):header+6],max(1,header-1))),last_row=len(grid),fields=FIELD_LABELS,mapping_fields=TEXT_FIELDS+NUM_FIELDS),400
            c=bp.db();cur=c.execute('INSERT INTO payroll_imports(run_id,data,warnings,source) VALUES(?,?,?,?)',(batch['id'],json.dumps(rows),json.dumps(warnings),u['filename']+' / '+sheet));iid=cur.lastrowid;c.commit();c.close()
            return redirect(url_for('billing.preview',iid=iid))
        return render_template('billing/map.html',batch=batch,upload=u,sheets=names,sheet=sheet,header=header,
            columns=[(i,str(grid[header-1][i] if i<len(grid[header-1]) else '')[:80]) for i in range(width)],
            sample=list(enumerate(grid[max(0,header-2):header+6],max(1,header-1))),last_row=len(grid),fields=FIELD_LABELS,mapping_fields=TEXT_FIELDS+NUM_FIELDS)
    except (ValueError,UnicodeError) as e:
        flash(str(e),'error')
        return redirect(url_for('billing.attendance_map',uid=uid,sheet=request.form.get('sheet',''),header=request.form.get('header',1))) if request.method=='POST' else redirect(url_for('billing.run',rid=batch['id']))
